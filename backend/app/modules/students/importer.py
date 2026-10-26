"""
Bulk import of students and guardians from a school's own spreadsheet (CSV or Excel).

Schools' sheets vary, so headers are matched loosely ("Adm No", "Admission
Number" and "REG NO" all mean the admission number), phone numbers are
normalised, and the class can be one column ("JSS 1A", "Primary 4 Gold") or
two (class + arm).

The whole file is checked first and every problem is reported with its row
number. `dry_run` stops there, so the console can show a preview. A real
run imports the valid rows and reports the rest. Re-importing the same file
is safe: existing admission numbers are skipped (or updated with
`mode="update"`).

Database work is a fixed handful of queries however many rows there are:
one lookup each for classes, existing students, existing guardians, their
links and enrolments, then a single flush.
"""

import csv
import enum
import io
import re
import uuid
from collections.abc import Iterable
from dataclasses import dataclass, field
from datetime import UTC, date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session, selectinload

from app.core.crypto import blind_index
from app.core.db import BULK_ROWS, refresh_statistics
from app.core.errors import Conflict, DomainError
from app.core.phone import InvalidPhone, normalize_phone
from app.modules.accounts.deps import StaffContext
from app.modules.audit.models import AuditAction
from app.modules.audit.service import record
from app.modules.schools.models import Arm, ClassLevel, Term

from .models import ConsentMethod, Enrollment, Gender, Guardian, GuardianStudent, Relationship, Student, StudentStatus

MAX_ROWS = 5000
PREVIEW_ROWS = 20


class ImportMode(enum.StrEnum):
    SKIP = "skip"  # leave students whose admission number already exists untouched
    UPDATE = "update"  # update their details, class and add new guardians


# --- Reading the file ----------------------------------------------------------------------------


def _norm(value: str) -> str:
    return re.sub(r"[^a-z0-9]", "", value.lower())


_STUDENT_ALIASES = {
    "admission_number": {
        "admissionnumber",
        "admissionno",
        "admno",
        "admission",
        "regno",
        "regnumber",
        "registrationnumber",
        "studentid",
        "studentno",
        "studentnumber",
        "idno",
    },
    "first_name": {"firstname", "givenname", "forename", "othernames", "othername"},
    "middle_name": {"middlename"},
    "last_name": {"lastname", "surname", "familyname"},
    "gender": {"gender", "sex"},
    "date_of_birth": {"dateofbirth", "dob", "birthdate", "birthday"},
    "class": {"class", "classlevel", "level", "grade", "form", "classname", "currentclass"},
    "arm": {"arm", "stream", "section", "classarm"},
}
# Guardian columns: a prefix (which guardian) + a suffix (which detail)
_GUARDIAN_PREFIXES = {
    "guardian": (1, None),
    "guardian1": (1, None),
    "parent": (1, None),
    "parent1": (1, None),
    "guardian2": (2, None),
    "parent2": (2, None),
    "father": (3, Relationship.FATHER),
    "mother": (4, Relationship.MOTHER),
}
_GUARDIAN_SUFFIXES = {
    "name": "name",
    "fullname": "name",
    "phone": "phone",
    "phonenumber": "phone",
    "mobile": "phone",
    "tel": "phone",
    "telephone": "phone",
    "gsm": "phone",
    "email": "email",
    "emailaddress": "email",
    "relationship": "relationship",
    "relation": "relationship",
}
# A lone "phone"/"email" column belongs to the first guardian
_LONE_GUARDIAN = {"phone": (1, "phone"), "phonenumber": (1, "phone"), "email": (1, "email")}


def _header_key(header: str) -> str | tuple[int, str] | None:
    key = _norm(header)
    for name, aliases in _STUDENT_ALIASES.items():
        if key in aliases:
            return name
    if key in _LONE_GUARDIAN:
        return _LONE_GUARDIAN[key]
    for prefix, (slot, _) in sorted(_GUARDIAN_PREFIXES.items(), key=lambda p: -len(p[0])):
        if key.startswith(prefix) and key[len(prefix) :] in _GUARDIAN_SUFFIXES:
            return slot, _GUARDIAN_SUFFIXES[key[len(prefix) :]]
        if key == prefix:  # a bare "Parent" column holds the name
            return slot, "name"
    return None


def _cell(value: Any) -> Any:
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # 8031234567.0 from Excel -> "8031234567"
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, int):
        return str(value)
    if isinstance(value, str):
        return value.strip()
    return value


def read_rows(filename: str, content: bytes) -> tuple[list[str], list[tuple[int, list[Any]]]]:
    """(headers, [(spreadsheet row number, cells)]) from a .csv or .xlsx file."""
    name = filename.lower()
    if name.endswith(".xlsx"):
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except Exception as e:
            raise DomainError("This file couldn't be opened as an Excel workbook.", field="file") from e
        sheet = workbook.worksheets[0]
        rows = [(n, [_cell(c) for c in row]) for n, row in enumerate(sheet.iter_rows(values_only=True), start=1)]
        workbook.close()
    elif name.endswith(".csv"):
        try:
            text = content.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = content.decode("cp1252", errors="replace")  # Excel's "CSV" on Windows
        rows = [(n, [_cell(c) for c in row]) for n, row in enumerate(csv.reader(io.StringIO(text)), start=1)]
    else:
        raise DomainError("Upload a .xlsx or .csv file.", field="file")
    rows = [(n, cells) for n, cells in rows if any(c not in (None, "") for c in cells)]
    if not rows:
        raise DomainError("The file is empty.", field="file")
    header_row, *data = rows
    if len(data) > MAX_ROWS:
        raise DomainError(f"Import at most {MAX_ROWS} students at a time.", field="file")
    return [str(h or "") for h in header_row[1]], data


# --- Parsing rows ----------------------------------------------------------------------------------


@dataclass
class RowError:
    row: int
    message: str
    column: str | None = None


@dataclass
class GuardianRow:
    full_name: str
    phone: str  # E.164
    phone_index: str
    email: str | None
    relationship: Relationship


@dataclass
class ParsedRow:
    row: int
    admission_number: str
    first_name: str
    middle_name: str
    last_name: str
    gender: Gender | None
    date_of_birth: date | None
    class_value: str
    arm_value: str
    guardians: list[GuardianRow]
    arm: Arm | None = None


_GENDERS = {
    "m": Gender.MALE,
    "male": Gender.MALE,
    "boy": Gender.MALE,
    "f": Gender.FEMALE,
    "female": Gender.FEMALE,
    "girl": Gender.FEMALE,
}
_RELATIONSHIPS = {r.value: r for r in Relationship} | {
    "mum": Relationship.MOTHER,
    "mom": Relationship.MOTHER,
    "dad": Relationship.FATHER,
    "parent": Relationship.GUARDIAN,
}


def _parse_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, date):
        return value
    text = str(value).strip()
    # Nigerian sheets write day first; ISO dates are also accepted
    for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%Y-%m-%d", "%d/%m/%y", "%d %b %Y", "%d %B %Y"):
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=UTC).date()
        except ValueError:
            continue
    raise ValueError(f"{text!r} isn't a date (use DD/MM/YYYY).")


def _split_other_names(first: str, middle: str) -> tuple[str, str]:
    """'Other names' columns often hold first and middle names together."""
    if not middle and " " in first:
        first, _, middle = first.partition(" ")
    return first, middle.strip()


def parse_rows(headers: list[str], data: list[tuple[int, list[Any]]]) -> tuple[list[ParsedRow], list[RowError]]:
    columns = [_header_key(h) for h in headers]
    found = {c for c in columns if isinstance(c, str)}
    missing = [
        label
        for key, label in (
            ("admission_number", "admission number"),
            ("first_name", "first name"),
            ("last_name", "surname"),
            ("class", "class"),
        )
        if key not in found
    ]
    if missing:
        raise DomainError(
            f"The file needs columns for: {', '.join(missing)}. Download the template to see the expected columns.",
            field="file",
            code="missing_columns",
        )
    parsed: list[ParsedRow] = []
    errors: list[RowError] = []
    seen_admission: dict[str, int] = {}
    for row_number, cells in data:
        values: dict[Any, Any] = {}
        for key, cell in zip(columns, cells, strict=False):
            if key is not None and cell not in (None, "") and key not in values:
                values[key] = cell
        row_errors: list[RowError] = []

        def need(
            key: str, label: str, n: int = row_number, v: dict[Any, Any] = values, errs: list[RowError] = row_errors
        ) -> str:
            value = str(v.get(key, "")).strip()
            if not value:
                errs.append(RowError(n, f"{label} is missing.", label))
            return value

        admission = need("admission_number", "Admission number")
        first, middle = _split_other_names(need("first_name", "First name"), str(values.get("middle_name", "")))
        last = need("last_name", "Surname")
        class_value = need("class", "Class")
        if admission and admission.lower() in seen_admission:
            row_errors.append(
                RowError(
                    row_number,
                    f"Admission number {admission} is also on row {seen_admission[admission.lower()]}.",
                    "Admission number",
                )
            )
        gender_raw = str(values.get("gender", "")).strip()
        gender = _GENDERS.get(gender_raw.lower()) if gender_raw else None
        if gender_raw and gender is None:
            row_errors.append(RowError(row_number, f"Gender {gender_raw!r} should be Male or Female.", "Gender"))
        try:
            dob = _parse_date(values.get("date_of_birth"))
        except ValueError as e:
            dob = None
            row_errors.append(RowError(row_number, str(e), "Date of birth"))
        guardians = _parse_guardians(row_number, values, row_errors)
        for text_value, limit, label in (
            (admission, 30, "Admission number"),
            (first, 60, "First name"),
            (last, 60, "Surname"),
            (middle, 60, "Middle name"),
        ):
            if len(text_value) > limit:
                row_errors.append(RowError(row_number, f"{label} is longer than {limit} characters.", label))
        if row_errors:
            errors.extend(row_errors)
            continue
        seen_admission[admission.lower()] = row_number
        parsed.append(
            ParsedRow(
                row_number,
                admission,
                first,
                middle,
                last,
                gender,
                dob,
                class_value,
                str(values.get("arm", "")).strip(),
                guardians,
            )
        )
    return parsed, errors


def _parse_guardians(row_number: int, values: dict[Any, Any], errors: list[RowError]) -> list[GuardianRow]:
    slots: dict[int, dict[str, str]] = {}
    for key, value in values.items():
        if isinstance(key, tuple):
            slot, detail = key
            slots.setdefault(slot, {})[detail] = str(value).strip()
    relationship_for_slot = dict(_GUARDIAN_PREFIXES.values())
    guardians: list[GuardianRow] = []
    seen: set[str] = set()
    for slot in sorted(slots):
        details = slots[slot]
        name, phone_raw = details.get("name", ""), details.get("phone", "")
        if not name and not phone_raw:
            continue
        label = {1: "Guardian", 2: "Guardian 2", 3: "Father", 4: "Mother"}[slot]
        if not phone_raw:
            errors.append(RowError(row_number, f"{label}'s phone number is missing.", f"{label} phone"))
            continue
        try:
            phone = normalize_phone(phone_raw)
        except InvalidPhone:
            errors.append(
                RowError(row_number, f"{label}'s phone {phone_raw!r} isn't a valid number.", f"{label} phone")
            )
            continue
        relationship = relationship_for_slot.get(slot) or _RELATIONSHIPS.get(
            details.get("relationship", "").lower(), Relationship.GUARDIAN
        )
        email = details.get("email") or None
        if email and ("@" not in email or len(email) > 255):
            errors.append(RowError(row_number, f"{label}'s email {email!r} isn't valid.", f"{label} email"))
            continue
        index = blind_index(phone)
        if index in seen:
            continue  # the same person listed twice on one row
        seen.add(index)
        guardians.append(GuardianRow(name or "Parent", phone, index, email, relationship))
    return guardians


# --- Matching classes ------------------------------------------------------------------------------


def _class_matcher(levels: list[ClassLevel]) -> Any:
    by_name = {_norm(level.name): level for level in levels}
    arms_by_level = {level.id: {_norm(a.name): a for a in level.arms if a.is_active} for level in levels}

    def match(class_value: str, arm_value: str) -> Arm | str:
        key = _norm(class_value)
        if arm_value:
            level = by_name.get(key)
            if level is None:
                return f"Class {class_value!r} doesn't exist. Add it under Classes, or fix the name."
            arm = arms_by_level[level.id].get(_norm(arm_value))
            return arm or f"{level.name} has no arm called {arm_value!r}."
        level = by_name.get(key)
        if level is not None:
            arms = list(arms_by_level[level.id].values())
            if len(arms) == 1:
                return arms[0]
            return f"{level.name} has several arms: add an Arm column or write the class like '{level.name} A'."
        # "JSS 1A" / "Primary 4 Gold": the longest class name that prefixes the value, then the arm
        for level_key in sorted(by_name, key=len, reverse=True):
            if key.startswith(level_key) and key != level_key:
                arm = arms_by_level[by_name[level_key].id].get(key[len(level_key) :])
                if arm is not None:
                    return arm
        return f"Class {class_value!r} doesn't exist. Add it under Classes, or fix the name."

    return match


# --- Running the import ----------------------------------------------------------------------------


@dataclass
class PreviewRow:
    row: int
    admission_number: str
    name: str
    class_name: str
    guardians: int
    action: str  # "create" | "update" | "skip"


@dataclass
class ImportReport:
    dry_run: bool
    total_rows: int
    students_created: int = 0
    students_updated: int = 0
    students_skipped: int = 0
    guardians_created: int = 0
    guardian_links_created: int = 0
    errors: list[RowError] = field(default_factory=list)
    preview: list[PreviewRow] = field(default_factory=list)


def import_students(
    session: Session,
    ctx: StaffContext,
    *,
    filename: str,
    content: bytes,
    term: Term,
    mode: ImportMode,
    consent_confirmed: bool,
    dry_run: bool,
) -> ImportReport:
    headers, data = read_rows(filename, content)
    rows, errors = parse_rows(headers, data)
    report = ImportReport(dry_run=dry_run, total_rows=len(data), errors=errors)
    if any(r.guardians for r in rows) and not consent_confirmed and not dry_run:
        raise DomainError(
            "Confirm that the guardians agreed to their details being stored.",
            field="consent_confirmed",
            code="consent_required",
        )

    # 1 query: classes and arms
    match = _class_matcher(
        list(
            session.scalars(
                select(ClassLevel).where(ClassLevel.is_active.is_(True)).options(selectinload(ClassLevel.arms))
            )
        )
    )
    placeable: list[ParsedRow] = []
    for row in rows:
        result = match(row.class_value, row.arm_value)
        if isinstance(result, str):
            report.errors.append(RowError(row.row, result, "Class"))
        else:
            row.arm = result
            placeable.append(row)

    # 1 query: students that already exist
    existing_students = {
        s.admission_number.lower(): s
        for s in session.scalars(
            select(Student).where(Student.admission_number.in_([r.admission_number for r in placeable]))
        )
    }
    # 1 query: guardians that already exist (siblings in the file share them too)
    guardian_rows = [g for r in placeable for g in r.guardians]
    guardians_by_index = {
        g.phone_index: g
        for g in session.scalars(
            select(Guardian).where(Guardian.phone_index.in_({g.phone_index for g in guardian_rows}))
        )
    }
    updating = mode == ImportMode.UPDATE
    touched_ids = [s.id for s in existing_students.values()] if updating else []
    # 1 query each: existing links and this term's enrolments for students being updated
    existing_links: set[tuple[uuid.UUID, uuid.UUID]] = set()
    has_primary: set[uuid.UUID] = set()
    enrollments: dict[uuid.UUID, Enrollment] = {}
    if touched_ids:
        for link in session.scalars(select(GuardianStudent).where(GuardianStudent.student_id.in_(touched_ids))):
            existing_links.add((link.guardian_id, link.student_id))
            if link.is_primary:
                has_primary.add(link.student_id)
        enrollments = {
            e.student_id: e
            for e in session.scalars(
                select(Enrollment).where(Enrollment.term_id == term.id, Enrollment.student_id.in_(touched_ids))
            )
        }

    now = datetime.now(UTC)
    new_objects: list[Any] = []
    for row in placeable:
        arm = row.arm
        if arm is None:  # placeable rows always have an arm
            continue
        student = existing_students.get(row.admission_number.lower())
        action = "create"
        if student is not None and not updating:
            report.students_skipped += 1
            _add_preview(report, row, arm, "skip")
            continue
        if student is None:
            student = Student(
                id=uuid.uuid4(),
                admission_number=row.admission_number,
                first_name=row.first_name,
                middle_name=row.middle_name,
                last_name=row.last_name,
                gender=row.gender,
                date_of_birth=row.date_of_birth,
                status=StudentStatus.ACTIVE,
            )
            new_objects.append(student)
            new_objects.append(
                Enrollment(student_id=student.id, term_id=term.id, arm_id=arm.id, class_level_id=arm.class_level_id)
            )
            report.students_created += 1
        else:
            action = "update"
            student.first_name, student.middle_name, student.last_name = (
                row.first_name,
                row.middle_name,
                row.last_name,
            )
            student.gender = row.gender or student.gender
            student.date_of_birth = row.date_of_birth or student.date_of_birth
            enrollment = enrollments.get(student.id)
            if enrollment is None:
                new_objects.append(
                    Enrollment(student_id=student.id, term_id=term.id, arm_id=arm.id, class_level_id=arm.class_level_id)
                )
            else:
                enrollment.arm_id, enrollment.class_level_id = arm.id, arm.class_level_id
            report.students_updated += 1
        for g in row.guardians:
            guardian = guardians_by_index.get(g.phone_index)
            if guardian is None:
                guardian = Guardian(
                    id=uuid.uuid4(),
                    full_name=g.full_name,
                    phone=g.phone,
                    phone_index=g.phone_index,
                    email=g.email,
                    consent_at=now,
                    consent_method=ConsentMethod.SCHOOL_ATTESTED,
                )
                guardians_by_index[g.phone_index] = guardian
                new_objects.append(guardian)
                report.guardians_created += 1
            if (guardian.id, student.id) in existing_links:
                continue
            primary = student.id not in has_primary
            has_primary.add(student.id)
            existing_links.add((guardian.id, student.id))
            new_objects.append(
                GuardianStudent(
                    guardian_id=guardian.id,
                    student_id=student.id,
                    relationship_to_student=g.relationship,
                    is_primary=primary,
                )
            )
            report.guardian_links_created += 1
        _add_preview(report, row, arm, action)

    report.errors.sort(key=lambda e: e.row)
    if dry_run:
        session.rollback()
        return report
    session.add_all(_insert_order(new_objects))
    try:
        session.flush()
    except IntegrityError as e:
        session.rollback()
        raise Conflict(
            "Some of these students or guardians were just added by someone else. Run the import again "
            "to pick up the rest."
        ) from e
    record(
        session,
        AuditAction.STUDENTS_IMPORTED,
        school_id=ctx.school_id,
        actor=ctx.actor,
        meta=ctx.meta,
        changes={
            "file": filename[:100],
            "term": term.id,
            "mode": mode,
            "rows": report.total_rows,
            "created": report.students_created,
            "updated": report.students_updated,
            "skipped": report.students_skipped,
            "errors": len(report.errors),
            "guardians_created": report.guardians_created,
        },
    )
    session.commit()
    if report.students_created >= BULK_ROWS:  # billing usually follows straight away
        refresh_statistics(session, ("students", "enrollments", "guardians", "guardian_students"))
    return report


def _insert_order(objects: Iterable[Any]) -> list[Any]:
    """Parents before children, so one flush can insert everything in batches."""
    rank = {Student: 0, Guardian: 0, Enrollment: 1, GuardianStudent: 1}
    return sorted(objects, key=lambda o: rank[type(o)])


def _add_preview(report: ImportReport, row: ParsedRow, arm: Arm, action: str) -> None:
    if len(report.preview) < PREVIEW_ROWS:
        name = " ".join(p for p in (row.first_name, row.middle_name, row.last_name) if p)
        class_name = f"{row.class_value} {row.arm_value}".strip()
        report.preview.append(PreviewRow(row.row, row.admission_number, name, class_name, len(row.guardians), action))


# --- Template ------------------------------------------------------------------------------------

TEMPLATE_HEADERS = [
    "Admission Number",
    "Surname",
    "First Name",
    "Middle Name",
    "Gender",
    "Date of Birth",
    "Class",
    "Arm",
    "Guardian Name",
    "Guardian Phone",
    "Guardian Email",
    "Guardian Relationship",
    "Guardian 2 Name",
    "Guardian 2 Phone",
    "Guardian 2 Email",
    "Guardian 2 Relationship",
]


def template_workbook(levels: list[ClassLevel]) -> bytes:
    workbook = Workbook()
    workbook.remove(workbook.worksheets[0])
    sheet = workbook.create_sheet("Students")
    sheet.append(TEMPLATE_HEADERS)
    level = levels[0] if levels else None
    sheet.append(
        [
            "ADM/2026/001",
            "Okafor",
            "Chiamaka",
            "Ada",
            "Female",
            "14/03/2015",
            level.name if level else "JSS 1",
            level.arms[0].name if level and level.arms else "A",
            "Ngozi Okafor",
            "08031234567",
            "",
            "Mother",
            "",
            "",
            "",
            "",
        ]
    )
    for column in sheet.columns:
        sheet.column_dimensions[column[0].column_letter].width = 20
    notes = workbook.create_sheet("How to fill this in")
    for line in (
        "One row per student. Delete the example row before importing.",
        "Required: Admission Number, Surname, First Name, Class.",
        "Class and Arm must match the classes set up in the console (e.g. Class 'JSS 1', Arm 'A'),",
        "or put both in the Class column, e.g. 'JSS 1A'.",
        "Gender: Male or Female. Date of Birth: DD/MM/YYYY. Both are optional.",
        "Phone numbers can be written as 0803..., 234803... or +234803....",
        "Brothers and sisters can share a guardian: use the same phone number.",
        "Relationship: Mother, Father, Guardian or Other.",
        "Available classes: "
        + (
            ", ".join(f"{lv.name} ({', '.join(a.name for a in lv.arms)})" for lv in levels)
            or "none yet; set them up under Classes first."
        ),
    ):
        notes.append([line])
    notes.column_dimensions["A"].width = 100
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
