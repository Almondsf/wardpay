"""
Report downloads: Excel (for working with the numbers) and PDF (for printing
and sharing). Both carry the school's name, what the report covers, and when
it was produced.

Money goes into Excel as exact decimal naira (never floats). Text that Excel
would treat as a formula (starting with =, +, -, @, tab or carriage return) is
prefixed with an apostrophe, so a name typed into a spreadsheet import can't
run as a formula on a bursar's computer.
"""

from collections.abc import Sequence
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.clock import LAGOS

from .fee_status import FeeCell, FeeStatusReport, PupilStatus
from .group import BranchFigures, GroupOverview
from .service import CollectedPayment, Collections, Debtor

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF = "application/pdf"
MONEY_FORMAT = "#,##0.00"
HEADER_FILL = PatternFill("solid", fgColor="0F5132")
METHOD_LABELS = {
    "virtual_account": "Transfer (student account)",
    "card": "Card",
    "checkout_transfer": "Bank transfer (online)",
    "ussd": "USSD",
    "cash": "Cash",
    "teller": "Bank teller",
    "pos": "POS",
    "bank_transfer": "Transfer (school account)",
}
_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")


def safe_text(value: str) -> str:
    """Neutralise spreadsheet formula injection."""
    return f"'{value}" if value.startswith(_FORMULA_START) else value


def naira(kobo: int) -> Decimal:
    return Decimal(kobo) / 100


def ngn(kobo: int) -> str:
    """For PDFs: the standard fonts have no naira sign."""
    sign = "-" if kobo < 0 else ""
    whole, rest = divmod(abs(kobo), 100)
    return f"{sign}NGN {whole:,}.{rest:02d}"


def _produced() -> str:
    return datetime.now(LAGOS).strftime("%d %b %Y, %H:%M")


def _day(value: date | None) -> str:
    return value.strftime("%d %b %Y") if value else ""


# --- Excel ------------------------------------------------------------------------------------------


def _sheet(
    ws: Worksheet,
    *,
    school: str,
    title: str,
    subtitle: str,
    headers: Sequence[tuple[str, int]],
    rows: Sequence[Sequence[Any]],
    money_columns: set[int],
    total: Sequence[Any] | None = None,
) -> None:
    ws.append([safe_text(school)])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([title])
    ws["A2"].font = Font(bold=True, size=12)
    ws.append([f"{subtitle}  ·  Produced {_produced()}"])
    ws.append([])
    header_row = ws.max_row + 1
    ws.append([h for h, _ in headers])
    for col, (_, width) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in rows:
        ws.append([safe_text(v) if isinstance(v, str) else v for v in row])
    if total is not None:
        ws.append(list(total))
        for cell in ws[ws.max_row]:
            cell.font = Font(bold=True)
    for col in money_columns:
        for (cell,) in ws.iter_rows(min_row=header_row + 1, min_col=col, max_col=col):
            cell.number_format = MONEY_FORMAT
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = f"A{header_row}:{get_column_letter(len(headers))}{header_row + len(rows)}"


def _save(wb: Workbook) -> bytes:
    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


DEBTOR_HEADERS = [
    ("Student", 26),
    ("Admission no.", 16),
    ("Class", 14),
    ("Owed (NGN)", 15),
    ("Oldest due", 13),
    ("Days overdue", 13),
    ("Parent", 26),
    ("Phone", 17),
    ("Reminders sent", 15),
    ("Last paid", 13),
    ("Note", 22),
]


def _debtor_row(d: Debtor) -> list[Any]:
    notes = [
        n for n in ("Reminders off" if d.reminders_off else "", "Left school" if d.status != "active" else "") if n
    ]
    return [
        d.full_name,
        d.admission_number,
        d.class_name,
        naira(d.owed_kobo),
        d.oldest_due,
        d.days_overdue,
        d.guardian_name,
        d.guardian_phone,
        d.reminders_sent,
        d.last_paid_at.astimezone(LAGOS).date() if d.last_paid_at else None,
        ", ".join(notes),
    ]


def debtors_xlsx(school: str, scope: str, debtors: list[Debtor]) -> bytes:
    wb = Workbook()
    ws = wb.create_sheet("Debtors", 0)
    wb.remove(wb.worksheets[1])  # the default empty sheet
    _sheet(
        ws,
        school=school,
        title="Debtors",
        subtitle=f"{scope}  ·  {len(debtors)} students",
        headers=DEBTOR_HEADERS,
        rows=[_debtor_row(d) for d in debtors],
        money_columns={4},
        total=["Total", "", "", naira(sum(d.owed_kobo for d in debtors))],
    )
    for (cell,) in ws.iter_rows(min_row=6, min_col=5, max_col=5):
        cell.number_format = "dd mmm yyyy"
    for (cell,) in ws.iter_rows(min_row=6, min_col=10, max_col=10):
        cell.number_format = "dd mmm yyyy"
    return _save(wb)


def collections_xlsx(school: str, summary: Collections, payments: list[CollectedPayment]) -> bytes:
    wb = Workbook()
    period = f"{_day(summary.start)} to {_day(summary.end)}"
    days = wb.create_sheet("By day", 0)
    wb.remove(wb.worksheets[1])  # the default empty sheet
    _sheet(
        days,
        school=school,
        title="Money received",
        subtitle=period,
        headers=[("Date", 14), ("Payments", 11), ("Net received (NGN)", 20)],
        rows=[[d.day, d.payments, naira(d.amount_kobo)] for d in summary.by_day],
        money_columns={3},
        total=["Total", summary.payments, naira(summary.net_kobo)],
    )
    for (cell,) in days.iter_rows(min_row=6, min_col=1, max_col=1):
        if isinstance(cell.value, date):
            cell.number_format = "ddd dd mmm yyyy"
    _sheet(
        wb.create_sheet("By method"),
        school=school,
        title="Money received by method",
        subtitle=period,
        headers=[("Method", 28), ("Payments", 11), ("Net received (NGN)", 20)],
        rows=[[METHOD_LABELS[m.method], m.payments, naira(m.amount_kobo)] for m in summary.by_method],
        money_columns={3},
        total=["Total", summary.payments, naira(summary.net_kobo)],
    )
    listing = wb.create_sheet("Payments")
    _sheet(
        listing,
        school=school,
        title="Payments and reversals",
        subtitle=period,
        headers=[
            ("Received", 18),
            ("Receipt no.", 12),
            ("Student", 26),
            ("Admission no.", 16),
            ("Class", 14),
            ("Method", 26),
            ("Reference", 22),
            ("Amount (NGN)", 15),
        ],
        rows=[
            [
                p.received_at.astimezone(LAGOS).replace(tzinfo=None),
                p.receipt_number,
                p.student_name,
                p.admission_number,
                p.class_name,
                ("Reversal: " if p.is_reversal else "") + METHOD_LABELS[p.method],
                p.reference,
                naira(p.amount_kobo),
            ]
            for p in payments
        ],
        money_columns={8},
    )
    for (cell,) in listing.iter_rows(min_row=6, min_col=1, max_col=1):
        if isinstance(cell.value, datetime):
            cell.number_format = "dd mmm yyyy hh:mm"
    return _save(wb)


# --- PDF --------------------------------------------------------------------------------------------


def _pdf(school: str, title: str, subtitle: str, blocks: list[Any]) -> bytes:
    styles = getSampleStyleSheet()
    buffer = BytesIO()

    def footer(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.grey)
        canvas.drawString(12 * mm, 8 * mm, f"{school} · {title} · produced {_produced()}")
        canvas.drawRightString(doc.pagesize[0] - 12 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()

    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=12 * mm,
        bottomMargin=14 * mm,
        title=title,
        author=school,
    )
    story = [
        Paragraph(escape(school), styles["Title"]),
        Paragraph(escape(title), styles["Heading2"]),
        Paragraph(escape(subtitle), styles["Normal"]),
        Spacer(1, 4 * mm),
        *blocks,
    ]
    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return buffer.getvalue()


def _table(
    header: list[str], rows: list[list[str]], *, widths: list[float], money_cols: set[int], total_row: bool = False
) -> Table:
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8, leading=10)
    money = ParagraphStyle("money", parent=cell, alignment=TA_RIGHT)
    bold_cell = ParagraphStyle("bold-cell", parent=cell, fontName="Helvetica-Bold")
    bold_money = ParagraphStyle("bold-money", parent=money, fontName="Helvetica-Bold")

    def style_for(row: int, col: int) -> ParagraphStyle:
        last = total_row and row == len(rows) - 1
        if col in money_cols:
            return bold_money if last else money
        return bold_cell if last else cell

    body = [[Paragraph(escape(v), style_for(r, c)) for c, v in enumerate(row)] for r, row in enumerate(rows)]
    table = Table([header, *body], colWidths=[w * mm for w in widths], repeatRows=1)
    style = [
        ("FONT", (0, 0), (-1, 0), "Helvetica-Bold", 8),
        ("FONT", (0, 1), (-1, -1), "Helvetica", 8),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f5132")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f2f5f3")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#d0d7d3")),
    ]
    style += [("ALIGN", (c, 0), (c, -1), "RIGHT") for c in money_cols]
    table.setStyle(TableStyle(style))
    return table


def debtors_pdf(school: str, scope: str, debtors: list[Debtor]) -> bytes:
    total = sum(d.owed_kobo for d in debtors)
    rows = [
        [
            d.full_name,
            d.admission_number,
            d.class_name,
            ngn(d.owed_kobo),
            f"{d.days_overdue} days" if d.days_overdue else "Not yet due",
            d.guardian_name,
            d.guardian_phone,
            "Reminders off" if d.reminders_off else str(d.reminders_sent),
        ]
        for d in debtors
    ]
    rows.append(["Total", "", "", ngn(total), "", "", "", ""])
    table = _table(
        ["Student", "Admission no.", "Class", "Owed", "Overdue", "Parent", "Phone", "Reminders"],
        rows,
        widths=[46, 30, 26, 30, 24, 50, 32, 25],
        money_cols={3},
        total_row=True,
    )
    return _pdf(school, "Debtors", f"{scope} · {len(debtors)} students owing {ngn(total)}", [table])


def collections_pdf(school: str, summary: Collections) -> bytes:
    styles = getSampleStyleSheet()
    period = f"{_day(summary.start)} to {_day(summary.end)}"
    totals = _table(
        ["Received", "Reversed", "Net received", "Payments"],
        [[ngn(summary.received_kobo), ngn(summary.reversed_kobo), ngn(summary.net_kobo), str(summary.payments)]],
        widths=[50, 50, 50, 30],
        money_cols={0, 1, 2},
    )
    methods = _table(
        ["Method", "Payments", "Net received"],
        [[METHOD_LABELS[m.method], str(m.payments), ngn(m.amount_kobo)] for m in summary.by_method],
        widths=[70, 30, 50],
        money_cols={2},
    )
    days = _table(
        ["Date", "Payments", "Net received"],
        [[d.day.strftime("%a %d %b %Y"), str(d.payments), ngn(d.amount_kobo)] for d in summary.by_day],
        widths=[50, 30, 50],
        money_cols={2},
    )
    return _pdf(
        school,
        "Money received",
        period,
        [
            totals,
            Spacer(1, 6 * mm),
            Paragraph("By method", styles["Heading3"]),
            methods,
            Spacer(1, 6 * mm),
            Paragraph("By day", styles["Heading3"]),
            days,
        ],
    )


PAID_FILL = PatternFill("solid", fgColor="D8EFE2")
PART_FILL = PatternFill("solid", fgColor="FCEBC7")
UNPAID_FILL = PatternFill("solid", fgColor="F8D7D3")
STATUS_FILLS = {"Paid": PAID_FILL, "Part paid": PART_FILL, "Unpaid": UNPAID_FILL}
FEE_STATUS_RULE = (
    "Money paid for particular fees is shown against those fees; the rest of a bill's payments "
    "cover tuition first, then its other fees in the order they appear on the bill."
)
PDF_STATUS_COLOURS = {"Paid": "#D8EFE2", "Part paid": "#FCEBC7", "Unpaid": "#F8D7D3"}
PDF_MAX_FEES = 9  # fee columns that fit across a landscape page; more share an "Other fees" column


def _class_columns(columns: list[str], pupils: list[PupilStatus]) -> list[str]:
    """The fees charged to these pupils, in bill order: each class lists only its own fees."""
    charged = {name for p in pupils for name, cell in p.fees.items() if cell.billed_kobo > 0}
    return [name for name in columns if name in charged]


def _by_class(pupils: list[PupilStatus]) -> dict[str, list[PupilStatus]]:
    groups: dict[str, list[PupilStatus]] = {}
    for p in pupils:  # already in class order
        groups.setdefault(p.class_name, []).append(p)
    return groups


def _fee_status_sheet(
    ws: Worksheet, school: str, report: FeeStatusReport, pupils: list[PupilStatus], scope: str
) -> None:
    columns = _class_columns(report.columns, pupils)
    headers: list[tuple[str, int]] = [("Student", 26), ("Admission no.", 16), ("Class", 14)]
    for name in columns:
        headers += [(name, 12), (f"{name}: paid (NGN)", 16), (f"{name}: left (NGN)", 16)]
    headers += [
        ("Billed this term (NGN)", 18),
        ("Paid this term (NGN)", 18),
        ("Left this term (NGN)", 18),
        ("Earlier debts (NGN)", 17),
        ("Total owed (NGN)", 16),
        ("Status", 12),
    ]
    rows = []
    for p in pupils:
        row: list[Any] = [p.full_name, p.admission_number, p.class_name]
        for name in columns:
            cell = p.fees.get(name)
            row += [cell.status, naira(cell.paid_kobo), naira(cell.left_kobo)] if cell else ["Not charged", None, None]
        row += [
            naira(p.term_billed_kobo),
            naira(p.term_paid_kobo),
            naira(p.term_left_kobo),
            naira(p.earlier_left_kobo),
            naira(p.total_left_kobo),
            p.status,
        ]
        rows.append(row)
    fee_cols = len(columns) * 3
    money = {4 + 3 * i + k for i in range(len(columns)) for k in (1, 2)} | set(range(4 + fee_cols, 4 + fee_cols + 5))
    total: list[Any] = ["Total", "", ""]
    for name in columns:
        cells = [p.fees[name] for p in pupils if name in p.fees]
        total += ["", naira(sum(c.paid_kobo for c in cells)), naira(sum(c.left_kobo for c in cells))]
    total += [
        naira(sum(p.term_billed_kobo for p in pupils)),
        naira(sum(p.term_paid_kobo for p in pupils)),
        naira(sum(p.term_left_kobo for p in pupils)),
        naira(sum(p.earlier_left_kobo for p in pupils)),
        naira(sum(p.total_left_kobo for p in pupils)),
        "",
    ]
    _sheet(
        ws,
        school=school,
        title=f"Fee status · {report.term_label}",
        subtitle=f"{scope}  ·  {len(pupils)} pupils",
        headers=headers,
        rows=rows,
        money_columns=money,
        total=total,
    )
    status_cols = [4 + 3 * i for i in range(len(columns))] + [len(headers)]
    first = ws.max_row - len(rows)  # the total row is last
    for col in status_cols:
        for (cell,) in ws.iter_rows(min_row=first, max_row=ws.max_row - 1, min_col=col, max_col=col):
            if cell.value in STATUS_FILLS:
                cell.fill = STATUS_FILLS[cell.value]
    ws.append([])
    ws.append([FEE_STATUS_RULE])


def fee_status_xlsx(school: str, report: FeeStatusReport) -> bytes:
    """One sheet for everyone selected; with all classes, also one sheet per class with that class's own fees."""
    wb = Workbook()
    _fee_status_sheet(wb.active, school, report, report.pupils, report.scope)
    wb.active.title = "All pupils" if report.scope == "All classes" else _sheet_title(report.scope)
    if report.scope == "All classes":
        for class_name, pupils in _by_class(report.pupils).items():
            _fee_status_sheet(wb.create_sheet(_sheet_title(class_name)), school, report, pupils, class_name)
    return _save(wb)


def _sheet_title(name: str) -> str:
    """Excel sheet names: at most 31 characters, none of the characters it forbids."""
    return "".join("-" if ch in "[]:*?/\\" else ch for ch in name)[:31] or "Class"


def _short(kobo: int) -> str:
    """Naira without the currency or whole-naira kobo, so a fee's state fits on one line ("55,000")."""
    return ngn(kobo).replace("NGN ", "").removesuffix(".00")


def _pdf_cell(cell: FeeCell | None) -> str:
    if cell is None or cell.billed_kobo <= 0:
        return "-"
    if cell.left_kobo <= 0:
        return "Paid"
    if cell.status == "Unpaid":
        return f"Unpaid {_short(cell.left_kobo)}"
    return f"{_short(cell.paid_kobo)} paid, {_short(cell.left_kobo)} left"


def _fee_status_class_table(columns: list[str], pupils: list[PupilStatus]) -> Table:
    """One class: its own fees across, each pupil's state on every fee, then the totals and status."""
    own = _class_columns(columns, pupils)
    shown = own if len(own) <= PDF_MAX_FEES else own[: PDF_MAX_FEES - 1]
    rest = own[len(shown) :]
    earlier = any(p.earlier_left_kobo for p in pupils)
    fee_count = len(shown) + (1 if rest else 0)
    header = [
        "Pupil",
        *shown,
        *(["Other fees"] if rest else []),
        "Paid this term",
        *(["Earlier debts"] if earlier else []),
        "Total owed",
        "Status",
    ]
    rows = []
    for p in pupils:
        cells = [_pdf_cell(p.fees.get(name)) for name in shown]
        if rest:
            others = [p.fees[n] for n in rest if n in p.fees]
            cells.append(_pdf_cell(FeeCell(sum(c.billed_kobo for c in others), sum(c.left_kobo for c in others))))
        rows.append(
            [
                p.full_name,
                *cells,
                ngn(p.term_paid_kobo),
                *([ngn(p.earlier_left_kobo)] if earlier else []),
                ngn(p.total_left_kobo),
                p.status,
            ]
        )
    rows.append(
        [
            "Total",
            *[""] * fee_count,
            ngn(sum(p.term_paid_kobo for p in pupils)),
            *([ngn(sum(p.earlier_left_kobo for p in pupils))] if earlier else []),
            ngn(sum(p.total_left_kobo for p in pupils)),
            "",
        ]
    )
    money_widths = [29] * (3 if earlier else 2)
    fixed = 44 + sum(money_widths) + 18
    each = (273 - fixed) / max(fee_count, 1)
    first_money = 1 + fee_count
    table = _table(
        header,
        rows,
        widths=[44, *[each] * fee_count, *money_widths, 18],
        money_cols=set(range(first_money, first_money + len(money_widths))),
        total_row=True,
    )
    status_col = len(header) - 1
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (status_col, r + 1), (status_col, r + 1), colors.HexColor(PDF_STATUS_COLOURS[p.status]))
                for r, p in enumerate(pupils)
                if p.status in PDF_STATUS_COLOURS
            ]
        )
    )
    return table


def fee_status_pdf(school: str, report: FeeStatusReport) -> bytes:
    """A table per class, each on its own page with only that class's fees, so it can be printed per class."""
    styles = getSampleStyleSheet()
    blocks: list[Any] = []
    for i, (class_name, pupils) in enumerate(_by_class(report.pupils).items()):
        if i:
            blocks.append(PageBreak())
        owing = sum(1 for p in pupils if p.total_left_kobo > 0)
        blocks += [
            Paragraph(escape(f"{class_name} · {len(pupils)} pupils, {owing} still owing"), styles["Heading3"]),
            _fee_status_class_table(report.columns, pupils),
            Spacer(1, 3 * mm),
            Paragraph(escape(f"Amounts in naira. {FEE_STATUS_RULE}"), styles["Italic"]),
        ]
    owing = sum(1 for p in report.pupils if p.total_left_kobo > 0)
    subtitle = f"{report.scope} · {len(report.pupils)} pupils, {owing} still owing"
    return _pdf(
        school, f"Fee status · {report.term_label}", subtitle, blocks or [Paragraph("No pupils.", styles["Normal"])]
    )


GROUP_HEADERS = [
    ("Branch", 28),
    ("Current term", 26),
    ("Pupils", 9),
    ("Billed this term (NGN)", 20),
    ("Collected (NGN)", 17),
    ("Owed this term (NGN)", 19),
    ("Owed from before (NGN)", 21),
    ("Pupils owing", 13),
    ("Credit (NGN)", 14),
    ("Received, 7 days (NGN)", 21),
    ("Received, 30 days (NGN)", 22),
    ("Awaiting approval (NGN)", 22),
]


def _group_row(b: BranchFigures) -> list[Any]:
    return [
        b.name,
        b.term_label or "No current term",
        b.pupils,
        naira(b.billed_kobo),
        naira(b.collected_kobo),
        naira(b.outstanding_kobo),
        naira(b.older_outstanding_kobo),
        b.students_owing,
        naira(b.credit_kobo),
        naira(b.received_7_days_kobo),
        naira(b.received_30_days_kobo),
        naira(b.pending_approval_kobo),
    ]


def group_xlsx(report: GroupOverview) -> bytes:
    """Head office: every branch side by side, with the group's totals."""
    wb = Workbook()
    ws = wb.create_sheet("Branches", 0)
    wb.remove(wb.worksheets[1])
    total = _group_row(report.totals)
    total[1] = ""
    _sheet(
        ws,
        school=report.group_name,
        title="All branches",
        subtitle=f"{len(report.branches)} branches  ·  as at {report.today:%d %b %Y}",
        headers=GROUP_HEADERS,
        rows=[_group_row(b) for b in report.branches],
        money_columns={4, 5, 6, 7, 9, 10, 11, 12},
        total=total,
    )
    return _save(wb)
