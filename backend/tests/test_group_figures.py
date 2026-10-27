"""Head-office figures: each branch side by side, the group's totals, and nothing from outside the group."""

import uuid
from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, update

from app.core.config import get_settings
from app.core.tenancy import scope_platform_wide, scope_to_school
from app.modules.accounts.models import AuthSession, Role, User
from app.modules.accounts.security import hash_password
from app.modules.audit.models import AuditAction, PlatformAuditLog
from app.modules.reports import group, service
from app.modules.schools.models import School, SchoolGroup, SchoolStatus
from tests.factories import (
    PASSWORD,
    QueryCounter,
    headers_for,
    make_calendar,
    make_levels,
    make_school,
    make_student,
    make_user,
)
from tests.test_fees import add_fees, generate
from tests.test_reports import DUE, TODAY, pay
from tests.test_reports import world as world  # the fixture: Grace Model College with bills and payments

FIGURES = (
    "billed_kobo",
    "collected_kobo",
    "outstanding_kobo",
    "older_outstanding_kobo",
    "students_owing",
    "credit_kobo",
    "pending_approval",
)


async def billed_school(db, client, name: str, pupils: int = 2) -> dict:
    school = make_school(db, name)
    bursar = make_user(db, school, Role.BURSAR)
    _, terms = make_calendar(school)
    levels = make_levels(school)
    add_fees(school, terms[0], levels)
    for i in range(pupils):
        make_student(school, terms[0], levels[0].arms[0], first=f"P{i}")
    db.commit()
    w = {"school": school, "terms": terms, "bursar": headers_for(db, bursar)}
    await generate(client, w, due_date=DUE)
    return w


@pytest.fixture
async def group_world(db, client, world, monkeypatch):
    monkeypatch.setattr(get_settings(), "enforce_mfa", False)
    lekki = await billed_school(db, client, "Lekki Campus")
    outsider = await billed_school(db, client, "Not In The Group")
    scope_platform_wide(db)
    g = SchoolGroup(id=uuid.uuid4(), name=f"Grace Group {uuid.uuid4().hex[:4]}")
    db.add(g)
    db.flush()
    for school_id in (world["school"].id, lekki["school"].id):
        db.get(School, school_id).group_id = g.id
    owner = User(
        id=uuid.uuid4(),
        group_id=g.id,
        role=Role.PROPRIETOR,
        full_name="Group Owner",
        email=f"owner-{uuid.uuid4().hex[:6]}@example.com",
        password_hash=hash_password(PASSWORD),
    )
    db.add(owner)
    db.commit()
    return {"group": g, "grace": world, "lekki": lekki, "outsider": outsider, "owner": owner}


def figures(db, gw, today=TODAY):
    scope_platform_wide(db)
    ids = tuple(db.scalars(select(School.id).where(School.group_id == gw["group"].id, School.status == "active")))
    return group.overview(db, gw["group"].id, ids, today=today)


def single(db, school_id, today=TODAY):
    db.info.pop("tenant_scope", None)
    scope_to_school(db, school_id)
    return service.overview(db, None, today=today)


async def test_each_branch_matches_its_own_dashboard_and_totals_add_up(db, group_world):
    gw = group_world
    report = figures(db, gw)
    assert [b.name for b in report.branches] == ["Grace Model College", "Lekki Campus"]
    grace = single(db, gw["grace"]["school"].id)
    row = report.branches[0]
    for name in FIGURES:
        assert getattr(row, name) == getattr(grace, name), name
    assert row.total_outstanding_kobo == grace.total_outstanding_kobo
    assert row.received_7_days_kobo == grace.received_7_days_kobo
    assert row.term_label == grace.term_label and row.pupils == 3
    lekki = single(db, gw["lekki"]["school"].id)
    assert report.branches[1].billed_kobo == lekki.billed_kobo > 0
    # Totals are the two branches, and nothing of the school outside the group
    for name in FIGURES:
        assert getattr(report.totals, name) == getattr(grace, name) + getattr(lekki, name), name
    # The group's daily takings over 30 days: Ada's 10k and Bayo's 50k on 1 October
    assert len(report.by_day) == 30 and report.by_day[-1].day == TODAY
    assert {d.day: d.amount_kobo for d in report.by_day if d.amount_kobo} == {date(2026, 10, 1): 60_000 * 100}


async def test_a_fixed_number_of_queries_however_many_branches(db, client, group_world, connection):
    gw = group_world
    with QueryCounter(connection) as small:
        figures(db, gw)
    for name in ("Ajah Campus", "Yaba Campus"):
        extra = await billed_school(db, client, name, pupils=4)
        scope_platform_wide(db)
        db.get(School, extra["school"].id).group_id = gw["group"].id
        db.commit()
    with QueryCounter(connection) as big:
        assert len(figures(db, gw).branches) == 4
    assert big.count == small.count


async def test_head_office_sees_it_branch_staff_and_platform_do_not(client, db, group_world):
    gw = group_world
    owner = headers_for(db, gw["owner"])
    response = await client.get("/api/v1/reports/group", headers=owner)
    assert response.status_code == 200, response.text
    body = response.json()
    assert [b["name"] for b in body["branches"]] == ["Grace Model College", "Lekki Campus"]
    assert body["totals"]["school_id"] is None and body["totals"]["name"] == "All branches"
    assert "Not In The Group" not in response.text
    # A branch's own staff and the platform admin get nothing
    assert (await client.get("/api/v1/reports/group", headers=gw["grace"]["bursar"])).status_code == 403
    admin = headers_for(db, make_user(db, None, Role.PLATFORM_ADMIN))
    assert (await client.get("/api/v1/reports/group", headers=admin)).status_code == 403
    # A suspended branch drops out at once
    scope_platform_wide(db)
    db.get(School, gw["lekki"]["school"].id).status = SchoolStatus.SUSPENDED
    db.commit()
    body = (await client.get("/api/v1/reports/group", headers=owner)).json()
    assert [b["name"] for b in body["branches"]] == ["Grace Model College"]


async def test_works_before_a_branch_is_chosen(client, db, group_world):
    gw = group_world
    owner = headers_for(db, gw["owner"])
    scope_platform_wide(db)
    db.execute(update(AuthSession).where(AuthSession.user_id == gw["owner"].id).values(branch_id=None))
    db.commit()
    assert (await client.get("/api/v1/students", headers=owner)).json()["error"]["code"] == "branch_required"
    assert (await client.get("/api/v1/reports/group", headers=owner)).status_code == 200


async def test_excel_download_is_recorded(client, db, group_world):
    gw = group_world
    await pay(client, gw["grace"], gw["grace"]["students"]["chidi"], 1_000, on=TODAY.isoformat())
    response = await client.get("/api/v1/reports/group/export", headers=headers_for(db, gw["owner"]))
    assert response.status_code == 200 and response.headers["cache-control"] == "no-store"
    ws = load_workbook(BytesIO(response.content)).active
    names = [row[0] for row in ws.iter_rows(min_row=6, values_only=True)]
    assert names == ["Grace Model College", "Lekki Campus", "All branches"]
    scope_platform_wide(db)
    entry = db.scalar(
        select(PlatformAuditLog).where(
            PlatformAuditLog.action == AuditAction.REPORT_EXPORTED.value,
            PlatformAuditLog.actor_user_id == gw["owner"].id,
        )
    )
    assert entry is not None and entry.changes["report"] == "group"
